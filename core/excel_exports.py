import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .models import User, Attendance, Invoice, Quote, Intervention, InventoryItem, Client, Contract


def _style_workbook(wb, title, headers, rows):
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    title_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="064E3B")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # Bandeau titre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell_title = ws.cell(row=1, column=1, value=f"ZECLEAN PRO — {title.upper()}")
    cell_title.font = title_font
    cell_title.fill = title_fill
    cell_title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    # En-têtes de colonnes
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws.row_dimensions[3].height = 24

    # Lignes de données
    for row_idx, r in enumerate(rows, start=4):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(r, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            if isinstance(val, (int, float)):
                c.number_format = '#,##0'
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Ajustement automatique de la largeur des colonnes
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws.cell(row=r_i, column=col_idx).value or ''))
            for r_i in range(3, len(rows) + 4)
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 5, 14), 45)


def _workbook_response(wb, filename):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def export_staff_excel(request):
    """Export Excel de la paie, des collaborateurs et des pointages."""
    wb = Workbook()
    headers = [
        "Matricule / ID", "Nom Complet", "Rôle", "Poste / Qualification",
        "Téléphone", "Date d'embauche", "Taux Horaire (FCFA)", "Salaire Mensuel (FCFA)", "Spécialités"
    ]
    staff = User.objects.filter(role__in=[User.ROLE_ADMIN, User.ROLE_SUPERVISOR, User.ROLE_AGENT]).order_by('role', 'last_name')
    rows = []
    for u in staff:
        rows.append([
            f"EMP-{u.id:03d}",
            u.get_full_name() or u.username,
            u.get_role_display(),
            u.job_title,
            u.phone,
            u.hire_date.strftime('%d/%m/%Y') if u.hire_date else '-',
            int(u.hourly_rate),
            int(u.monthly_salary),
            u.specialties,
        ])
    _style_workbook(wb, "Personnel & Paie (FCFA)", headers, rows)

    # Deuxième onglet : Pointages
    ws2 = wb.create_sheet(title="Feuille de Pointage")
    ws2.append(["Date", "Collaborateur", "Site affecté", "Heure Arrivée", "Heure Départ", "Heures Travaillées", "Statut", "Remarques"])
    for att in Attendance.objects.select_related('agent', 'site').all():
        ws2.append([
            att.date.strftime('%d/%m/%Y'),
            att.agent.get_full_name(),
            att.site.name if att.site else '-',
            att.check_in.strftime('%H:%M') if att.check_in else '-',
            att.check_out.strftime('%H:%M') if att.check_out else '-',
            float(att.hours_worked),
            att.get_status_display(),
            att.notes,
        ])

    return _workbook_response(wb, "ZeCleanPro_Paie_Pointage.xlsx")


@login_required
def export_billing_excel(request):
    """Export Excel des factures, paiements et devis en FCFA."""
    wb = Workbook()
    headers = [
        "N° Facture", "Date Émission", "Échéance", "Client",
        "Objet / Période", "Montant HT (FCFA)", "TVA 18% (FCFA)",
        "Total TTC (FCFA)", "Déjà Réglé (FCFA)", "Reste à Payer (FCFA)", "Statut", "Mode de règlement"
    ]
    invoices = Invoice.objects.select_related('client').all()
    rows = []
    for inv in invoices:
        rows.append([
            inv.number,
            inv.issue_date.strftime('%d/%m/%Y'),
            inv.due_date.strftime('%d/%m/%Y'),
            inv.client.name,
            inv.period_label,
            int(inv.amount_ht),
            int(inv.tax_amount),
            int(inv.amount_ttc),
            int(inv.amount_paid),
            int(inv.balance_due),
            inv.get_status_display(),
            inv.get_payment_method_display(),
        ])
    _style_workbook(wb, "Facturation & Recouvrement", headers, rows)

    # Deuxième onglet : Devis commerciaux
    ws2 = wb.create_sheet(title="Devis Commerciaux")
    ws2.append(["N° Devis", "Date", "Client", "Site concerné", "Surface (m²)", "Montant HT (FCFA)", "Montant TTC (FCFA)", "Statut"])
    for q in Quote.objects.select_related('client').all():
        ws2.append([
            q.number,
            q.issue_date.strftime('%d/%m/%Y'),
            q.client.name,
            q.site_name,
            q.surface_m2,
            int(q.amount_ht),
            int(q.amount_ttc),
            q.get_status_display(),
        ])

    return _workbook_response(wb, "ZeCleanPro_Facturation_FCFA.xlsx")


@login_required
def export_planning_excel(request):
    """Export Excel du planning des interventions."""
    wb = Workbook()
    headers = [
        "Code", "Date Prévue", "Horaires", "Titre de la mission",
        "Type de Prestation", "Client", "Site", "Surface (m²)",
        "Équipe Assignée", "Priorité", "Statut", "Checklist (%)"
    ]
    interventions = Intervention.objects.select_related('site', 'site__client', 'team').all()
    rows = []
    for inter in interventions:
        rows.append([
            inter.code,
            inter.scheduled_date.strftime('%d/%m/%Y'),
            f"{inter.start_time.strftime('%H:%M')} - {inter.end_time.strftime('%H:%M')}",
            inter.title,
            inter.get_intervention_type_display(),
            inter.site.client.name,
            inter.site.name,
            inter.site.surface_m2,
            inter.team.name if inter.team else "Agents individuels",
            inter.get_priority_display(),
            inter.get_status_display(),
            inter.checklist_progress,
        ])
    _style_workbook(wb, "Planning des Interventions", headers, rows)
    return _workbook_response(wb, "ZeCleanPro_Planning_Interventions.xlsx")


@login_required
def export_inventory_excel(request):
    """Export Excel de l'état des stocks et équipements."""
    wb = Workbook()
    headers = [
        "Code SKU", "Désignation", "Catégorie", "Quantité en Stock",
        "Unité", "Seuil d'alerte", "Coût Unitaire HT (FCFA)",
        "Valeur Totale Stock (FCFA)", "Fournisseur", "Alerte"
    ]
    items = InventoryItem.objects.all()
    rows = []
    for it in items:
        rows.append([
            it.sku,
            it.name,
            it.get_category_display(),
            float(it.quantity),
            it.unit,
            float(it.min_threshold),
            int(it.unit_cost),
            int(it.total_value),
            it.supplier,
            "COMMANDER (STOCK BAS)" if it.is_low_stock else "OK",
        ])
    _style_workbook(wb, "Inventaire Stocks & Matériel", headers, rows)
    return _workbook_response(wb, "ZeCleanPro_Stocks_Materiel.xlsx")


@login_required
def export_clients_excel(request):
    """Export Excel du portefeuille clients et contrats."""
    wb = Workbook()
    headers = [
        "Client / Raison Sociale", "Catégorie", "Interlocuteur",
        "Téléphone", "Email", "Adresse", "Ville", "NINEA / RCCM", "Nb Sites", "Nb Contrats"
    ]
    clients = Client.objects.prefetch_related('sites', 'contracts').all()
    rows = []
    for c in clients:
        rows.append([
            c.name,
            c.get_client_type_display(),
            c.contact_person,
            c.phone,
            c.email,
            c.address,
            c.city,
            c.tax_id,
            c.sites.count(),
            c.contracts.count(),
        ])
    _style_workbook(wb, "Portefeuille Clients", headers, rows)

    ws2 = wb.create_sheet(title="Contrats Récurrents")
    ws2.append(["Référence", "Client", "Site", "Prestation", "Fréquence", "Forfait Mensuel HT (FCFA)", "Statut"])
    for ctr in Contract.objects.select_related('client', 'site').all():
        ws2.append([
            ctr.reference,
            ctr.client.name,
            ctr.site.name,
            ctr.title,
            ctr.get_frequency_display(),
            int(ctr.monthly_amount_ht),
            ctr.get_status_display(),
        ])

    return _workbook_response(wb, "ZeCleanPro_Clients_Contrats.xlsx")
